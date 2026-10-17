"""Create student-personalized reflection and evidence-backed technical report."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK

ROOT=Path(__file__).resolve().parents[1]
PROJECT=ROOT/'ITAI2373-NewsBot-Midterm'
DOCS=PROJECT/'documents'
DOCS.mkdir(exist_ok=True)
URL='https://github.com/MarineMaaan/ITAI2373-NewsBot-Midterm'
results=json.loads((PROJECT/'outputs'/'results.json').read_text(encoding='utf-8'))
winner=next(r for r in results['test_metrics'] if r['model']==results['selected_model'])
research=results['research']

def setup():
    d=Document()
    s=d.sections[0]
    s.page_width=Inches(8.5); s.page_height=Inches(11)
    s.top_margin=Inches(.68); s.bottom_margin=Inches(.68)
    s.left_margin=Inches(.8); s.right_margin=Inches(.8)
    normal=d.styles['Normal']
    normal.font.name='Calibri'; normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(7)
    normal.paragraph_format.line_spacing=1.05
    for name,size in [('Title',27),('Subtitle',11),('Heading 1',15),('Heading 2',12)]:
        st=d.styles[name]; st.font.name='Calibri'; st.font.size=Pt(size)
        st.font.color.rgb=RGBColor.from_string('173B55')
        st.paragraph_format.space_before=Pt(9); st.paragraph_format.space_after=Pt(5)
        for border in st.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    d.core_properties.author='Benjamin LaCount II'
    d.core_properties.subject='ITAI 2373 NewsBot solo midterm'
    return d

def link(d,url):
    p=d.add_paragraph()
    h=OxmlElement('w:hyperlink')
    h.set(qn('r:id'),d.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True))
    run=OxmlElement('w:r'); props=OxmlElement('w:rPr'); color=OxmlElement('w:color'); color.set(qn('w:val'),'147D92'); props.append(color)
    run.append(props); t=OxmlElement('w:t'); t.text=url; run.append(t); h.append(run); p._p.append(h)

def title(d,heading,subtitle):
    d.add_paragraph(heading,'Title'); d.add_paragraph(subtitle,'Subtitle')

def section(d,heading,text):
    d.add_heading(heading,level=1)
    for p in text.split('\n\n'): d.add_paragraph(p)

def table(d,headers,rows):
    t=d.add_table(rows=1,cols=len(headers)); t.style='Light Shading Accent 1'
    for cell,header in zip(t.rows[0].cells,headers): cell.text=header
    pr=t.rows[0]._tr.get_or_add_trPr(); repeat=OxmlElement('w:tblHeader'); pr.append(repeat)
    for row in rows:
        for cell,value in zip(t.add_row().cells,row): cell.text=str(value)
    return t

# Reflection is exactly two deliberately composed pages, grounded in the student's notes.
d=setup()
title(d,'Reflective Journal','Benjamin LaCount II  |  ITAI 2373  |  October 17, 2026  |  Solo project')
link(d,URL)
section(d,'Independent work and preparation',
"""I completed this project as a solo assignment. There was no group name or number, and no video demonstration was included. In place of team role allocation, the work required connecting planning, implementation, analysis and documentation into one coherent deliverable. This made consistency especially important: a decision about the dataset or preprocessing had consequences throughout the notebook.

I took more time than most to prepare for the assignment. That preparation helped clarify the requirements, but it also showed me that I need a more efficient way to begin. In future projects, I want to turn the instructions into a short requirements checklist, confirm the data format early, and establish a small working pipeline before expanding the analysis. Preparation should reduce uncertainty without delaying the first useful result.""")
section(d,'Learning from preprocessing',
"""My strongest lesson was how much preprocessing influences the quality of the final output. Cleaning text is a modeling decision, not merely a preliminary chore. Removing punctuation, capitalization, stop words or numbers can simplify a classifier while also removing information needed by another NLP task.

The NewsBot design addresses that tension by preserving normalized article text for sentiment, parsing and entity recognition, while using tokenized, lemmatized content for TF-IDF classification. The BBC source is largely lowercased, so preserving its text cannot recover capitalization that is already missing. Recognizing that limit is part of interpreting the results responsibly.""")
section(d,'Technical integration challenges',
"""The central integration challenge was making several NLP techniques work together without giving them identical inputs or overstating what they measure. A category prediction, a sentiment score and an extracted relationship answer different questions. An entity co-occurrence edge shows that two names appeared together; it does not establish a business relationship.

Evaluation creates another important boundary. Vocabulary learning and feature selection belong inside training cross-validation, while reserved test articles measure generalization. Duplicate removal must happen before splitting. These choices connect data preparation to the credibility of the final model comparison, and they provide a more useful lesson than a high accuracy score alone.""")
d.add_page_break()
title(d,'From news analysis to studio insight','Reflective Journal  |  Benjamin LaCount II  |  Page 2')
section(d,'Business value assessment',
"""The application I care about is reading and categorizing game reviews for products my studio produces. A review-monitoring system could help identify recurring concerns about bugs, performance, gameplay and value, then direct attention to the most useful feedback. The practical benefit would be a clearer view of player concerns and a more efficient review process.

NewsBot is a foundation for that idea rather than proof that it already solves it. Its BBC labels describe news topics, not review themes. Whole-article sentiment can also miss whether a reviewer likes a game but dislikes one feature. A studio version would need appropriately sourced review data, new labels and evaluation across different games and review sources.""")
section(d,'Individual contribution and AI assistance',
"""I used AI to assist with executing the work while preserving the thought, planning and final quality assurance as my own responsibilities. That assistance included implementation, debugging, evaluation support and drafting documentation. I supplied the personal reflection themes and the studio use case. I remain responsible for understanding the pipeline, checking the outputs and making the final submission decisions.

Working solo does not remove the need for accountability. The repository makes the process inspectable through the notebook, reusable code, evaluation outputs and written explanations. Clear attribution of assistance matters because the project should demonstrate understanding and judgment as well as the ability to produce a working artifact.""")
section(d,'Future enhancements and professional development',
"""My next steps would be to prepare more efficiently, build a small review annotation set, and compare useful review categories before adding more model complexity. Human review of uncertain predictions should remain part of the workflow. Entity extraction and emotion analysis also need their own labeled evaluation data rather than relying on the topic classifier's results.

The bonus experiments broaden the project through an interactive dashboard, a controlled bilingual comparison, custom domain NER training and a measured feature ablation. Their small or synthetic evaluation sets have limits. For professional development, the key habit is to state those limits clearly, connect results to a real user need, and make the evidence reproducible. That approach is directly relevant to building analytical tools for my studio.""")
d.save(DOCS/'NewsBot_Reflection_Benjamin_LaCount_II.docx')

d=setup()
title(d,'NewsBot Intelligence System','Technical Report  |  Benjamin LaCount II  |  ITAI 2373  |  October 17, 2026')
link(d,URL)
section(d,'Purpose and scope',
"""NewsBot is a CPU-based Google Colab prototype for assisted news routing and media monitoring. It combines eight NLP modules into one reusable analysis function and an interactive Gradio interface. The student works solo; no team name, team number or video is used. The notebook requires a user to download and upload the labeled BBC training CSV. Dataset approval was confirmed by the student. Raw dataset files are excluded from the public repository.

The intended operator is an editor or analyst who needs topic routing, entity leads and tone summaries. Benjamin's studio review-monitoring use case motivates future adaptation, but news-topic performance does not establish game-review performance.""")
section(d,'Dataset and evaluation design',
f"""The retained dataset contains {results['dataset']['retained_rows']} substantial articles after removing {results['dataset']['duplicates_removed']} canonical duplicates and {results['dataset']['missing_removed']} rows with missing values. It uses BBC's business, entertainment, politics, sport and tech labels. Health is outside the trained label space. There are no reliable publication timestamps in the supplied CSV.

Stratified partition sizes are {results['split_sizes']['train']} training, {results['split_sizes']['validation']} validation and {results['split_sizes']['test']} test articles. Three-fold training CV optimizes n-grams, feature selection and model settings. TF-IDF vocabulary and chi-squared selection are fitted within folds. Validation determines a review-margin threshold; test data does not select the classifier. Exact duplicates are removed before splitting, although near duplicates may remain.""")
table(d,['Model','CV-selected?','Test accuracy','Test macro-F1'],[
    [r['model'],'Yes' if r['model']==results['selected_model'] else 'No',f"{r['accuracy']:.4f}",f"{r['macro_f1']:.4f}"] for r in results['test_metrics']])
section(d,'Main result',f"""Training CV selected {results['selected_model']}. Its held-out accuracy is {winner['accuracy']:.4f} and macro-F1 is {winner['macro_f1']:.4f}. The notebook includes per-class precision, recall, F1, support, a confusion matrix and an error table. These scores measure topic classification on this split of BBC data. They do not measure sentiment quality, entity accuracy or readiness for a live editorial workflow.""")
d.add_page_break()
title(d,'Integrated NLP and bonus evaluation','Technical Report  |  Benjamin LaCount II  |  Page 2')
table(d,['Module','Implementation and interpretation'],[
 ['1. Application','Editorial routing, human review and a future studio review-monitoring use case.'],
 ['2. Preprocessing','Markup/Unicode/URL cleanup; tokenization, stop words and lemmatization; raw text retained.'],
 ['3. TF-IDF','Fold-safe unigram/bigram search, chi-squared selection, category terms and model weights.'],
 ['4. POS','Per-article POS proportions, category heatmap and grammatical bigrams.'],
 ['5. Syntax','Dependency features and heuristic agent/action/patient relations, including passive voice and negation.'],
 ['6. Tone','VADER sentiment and transparent four-emotion cue lexicons; no gold tone labels.'],
 ['7. Classification','Naive Bayes, logistic regression and linear SVM with CV and reserved-test evaluation.'],
 ['8. NER','PERSON, ORG, GPE, DATE, MONEY; frequencies, category patterns and co-occurrence map.'],
])
section(d,'Bonus experiments',
f"""Interactive dashboard: Gradio accepts a new English article and returns category, tone, entities, syntax, custom entities and a review flag. Two original articles demonstrate the complete flow. The validation margin is not a calibrated probability or an out-of-domain detector.

Advanced analysis: 25 balanced original English/Spanish pairs compare English input, direct Spanish input and a fixed glossary aid. The passages intentionally contain domain keywords, making this an optimistic controlled language-sensitivity probe rather than a general multilingual benchmark.

Custom model: a small spaCy NER model learns NEWS_OUTLET and TECH_PRODUCT using 96 synthetic training, 12 development and 12 test sentences. Name pairs and sentence frames are disjoint. Development F1 selects the epoch; exact-span test precision is {results['custom_ner']['test_precision']:.3f}, recall {results['custom_ner']['test_recall']:.3f}, and F1 {results['custom_ner']['test_f1']:.3f}. The no-custom-label baseline has F1 0.000. Real article performance remains unmeasured.""")
d.add_page_break()
title(d,'Findings, limits and reproducibility','Technical Report  |  Benjamin LaCount II  |  Page 3')
section(d,'Research extension',
f"""The project-specific hypothesis is that adding grammar features to TF-IDF improves topic classification. A fixed logistic regression baseline and an otherwise matched hybrid model are evaluated on the same test articles. Baseline macro-F1 is {research['baseline_f1']:.4f}; hybrid macro-F1 is {research['hybrid_f1']:.4f}. The 95% stratified paired-bootstrap interval for hybrid minus baseline is [{research['difference_ci95'][0]:.4f}, {research['difference_ci95'][1]:.4f}] over 500 resamples.

An entity-word masking stress test produces baseline macro-F1 {research['baseline_masked_f1']:.4f} and hybrid macro-F1 {research['hybrid_masked_f1']:.4f}. Original syntax is retained in this lexical stress test. The approach is an original project experiment, not a claim of academic novelty. One split and a bootstrap conditional on fitted models cannot establish a universal benefit.""")
section(d,'Business recommendations and limitations',
"""Use NewsBot for assisted sorting, with a human examining uncertain cases and checking extracted entities in the source. The term and POS charts can guide exploration; they do not establish causal differences between categories. VADER measures lexical tone, and the emotion lexicons are unvalidated proxies. Co-occurrence edges are leads, not verified relationships. Lowercased BBC text weakens pretrained NER.

For studio reviews, create licensed review data and labels reflecting bugs, performance, gameplay and value. Evaluate across games, sources and time. Annotate aspect sentiment and domain entities separately. Measure operator time and error costs in a pilot before claiming productivity gains. Repeat experiments with independently annotated, representative material.""")
section(d,'Reproduction and evidence',
f"""Open NewsBot_Intelligence_System.ipynb in Colab, run setup, and upload BBC News Train.csv when prompted. Run the remaining cells in order. The notebook embeds reusable source and exports result tables, figures and JSON. Local validation uses the same notebook in a fresh CPU kernel with an explicitly supplied CSV path. This is local execution evidence; a signed-in Colab browser session is not claimed.

Machine-readable outputs record the dataset SHA-256, category counts, split manifest, library versions and runtime. The separately supplied reflective journal is two pages and records the student's preparation, learning, studio use case and AI assistance. The four non-video bonus categories are implemented; points and the stated extra-credit cap remain the instructor's decision.""")
section(d,'Sources and contribution',
"""BBC News Classification: kaggle.com/competitions/learn-ai-bbc/data. Technical references: scikit-learn.org/stable/common_pitfalls.html; spacy.io/usage/linguistic-features; spacy.io/usage/training; nltk.org/api/nltk.sentiment.vader.html; gradio.app/docs/gradio/interface. VADER: Hutto and Gilbert (2014), Proceedings of ICWSM, doi:10.1609/icwsm.v8i1.14550. Full clickable references are in the notebook and README.

Benjamin LaCount II is the sole student contributor. AI assisted execution, code, debugging, evaluation and documentation. The student supplied the reflection themes and retains thought, planning, understanding and final quality-assurance responsibilities. The course starter supplied the original module sequence.""")
d.save(DOCS/'MT_Report_Benjamin_LaCount_II_ITAI2373.docx')
print('Created two-page reflection and three-page technical report sources.')
